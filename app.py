import streamlit as st
import google.generativeai as genai
from docx import Document
from io import BytesIO

# -------------------------
# PAGE CONFIG
# -------------------------
st.set_page_config(
    page_title="AI Deal Brief Generator",
    layout="wide"
)

st.title("AI Deal Brief Generator")

# -------------------------
# GEMINI CONFIG
# -------------------------
genai.configure(
    api_key=st.secrets["GEMINI_API_KEY"]
)

model = genai.GenerativeModel("gemini-1.5-flash")

# -------------------------
# FORM
# -------------------------
with st.form("deal_form"):

    deal_name = st.text_input("Deal Name")

    client_name = st.text_input("Client")

    stage = st.selectbox(
        "Deal Stage",
        [
            "Discovery",
            "Qualified",
            "Proposal",
            "Negotiation",
            "Closed Won"
        ]
    )

    contacts = st.text_input("Client Contacts")

    amount = st.text_input("Deal Amount")

    date = st.date_input("Expected Close Date")

    tech_stack = st.text_area(
        "Tech Stack"
    )

    problem_statement = st.text_area(
        "Problem Statement / Notes",
        height=200
    )

    submitted = st.form_submit_button("Generate Deal Brief")

# -------------------------
# GENERATE
# -------------------------
if submitted:

    prompt = f"""
    Generate a professional enterprise deal brief.

    Deal Name: {deal_name}
    Client: {client_name}
    Stage: {stage}
    Contacts: {contacts}
    Amount: {amount}
    Date: {date}
    Tech Stack: {tech_stack}

    Problem Statement:
    {problem_statement}

    Include:
    - Executive Summary
    - GenAI Use Cases
    - Risks
    - Competitive Positioning
    - Resource Requirements
    - Next Steps
    - Qualification Questions
    """

    with st.spinner("Generating..."):

        response = model.generate_content(prompt)

        output = response.text

    st.subheader("Generated Deal Brief")

    st.markdown(output)

    st.code(output)

    # -------------------------
    # WORD DOC
    # -------------------------
    doc = Document()

    doc.add_heading("AI Deal Brief", level=1)

    doc.add_paragraph(output)

    buffer = BytesIO()

    doc.save(buffer)

    buffer.seek(0)

    st.download_button(
        label="Download Word Document",
        data=buffer,
        file_name=f"{deal_name}_Deal_Brief.docx",
        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
